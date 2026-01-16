"""Word报告生成服务"""

from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import os
from typing import Dict, Any, List, Optional
import logging

from .qwen_client import QwenClient
from .chart_generator import ChartGenerator

logger = logging.getLogger(__name__)

class ReportGenerator:
    """Word报告生成器"""

    def __init__(self, report_dir: str):
        """
        初始化报告生成器

        Args:
            report_dir: 单个报告的目录路径 (/root/librechat_user_data/{user_id}/reports/{report_id}/)
        """
        self.report_dir = report_dir
        os.makedirs(report_dir, exist_ok=True)

        self.qwen_client = QwenClient()
        # 在报告目录下创建charts子文件夹
        self.chart_generator = ChartGenerator(os.path.join(report_dir, 'charts'))
    
    def _set_chinese_font(self, run, font_name='宋体', font_size=12):
        """设置中文字体"""
        run.font.name = font_name
        run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
        run.font.size = Pt(font_size)
    
    def _add_heading(self, doc: Document, text: str, level: int = 1):
        """添加标题"""
        heading = doc.add_heading(text, level=level)
        for run in heading.runs:
            self._set_chinese_font(run, '黑体', 14 if level == 1 else 12)
    
    def _add_paragraph(self, doc: Document, text: str, bold: bool = False):
        """添加段落"""
        para = doc.add_paragraph()
        run = para.add_run(text)
        self._set_chinese_font(run, '宋体', 12)
        run.bold = bold
    
    def _add_table(
        self, 
        doc: Document, 
        headers: List[str], 
        rows: List[List[str]],
        title: Optional[str] = None
    ):
        """添加表格"""
        if title:
            self._add_paragraph(doc, title, bold=True)
        
        table = doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # 表头
        header_cells = table.rows[0].cells
        for i, header in enumerate(headers):
            header_cells[i].text = header
            for para in header_cells[i].paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.bold = True
                    self._set_chinese_font(run, '宋体', 10)
        
        # 数据行
        for row_idx, row_data in enumerate(rows):
            row_cells = table.rows[row_idx + 1].cells
            for col_idx, cell_data in enumerate(row_data):
                row_cells[col_idx].text = str(cell_data)
                for para in row_cells[col_idx].paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in para.runs:
                        self._set_chinese_font(run, '宋体', 10)
        
        doc.add_paragraph()  # 表格后空一行
    
    def _add_image(self, doc: Document, image_path: str, caption: str = None, width: float = 5.5):
        """添加图片"""
        if os.path.exists(image_path):
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run()
            run.add_picture(image_path, width=Inches(width))
            
            if caption:
                caption_para = doc.add_paragraph()
                caption_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = caption_para.add_run(caption)
                self._set_chinese_font(run, '宋体', 10)
                run.italic = True
    
    def generate_report(
        self,
        model_result: Dict[str, Any],
        preprocessing_record: Dict[str, Any],
        test_results: Dict[str, Any],
        descriptive_stats: Dict[str, Any],
        data_source_info: Dict[str, str],
        model_type: str
    ) -> str:
        """
        生成完整学术报告
        
        Args:
            model_result: 模型结果
            preprocessing_record: 预处理记录
            test_results: 统计检验结果
            descriptive_stats: 描述性统计 (支持单变量和多变量)
            data_source_info: 数据源信息
            model_type: 模型类型
        
        Returns:
            报告文件路径
        """
        doc = Document()
        
        # ========== 第1章：数据说明与处理 ==========
        self._add_heading(doc, '1 数据说明与处理', level=1)
        
        # 1.1 数据来源
        self._add_heading(doc, '1.1 数据来源', level=2)
        source_rows = [
            ['数据名称', data_source_info.get('name', '-')],
            ['数据来源', data_source_info.get('source', '-')],
            ['获取时间', data_source_info.get('fetch_time', '-')],
            ['时间范围', data_source_info.get('date_range', '-')],
            ['原始记录数', str(data_source_info.get('original_count', '-'))]
        ]
        self._add_table(doc, ['项目', '内容'], source_rows, '表1-1 数据来源信息')
        
        # 1.2 数据预处理
        self._add_heading(doc, '1.2 数据预处理', level=2)
        preprocess_rows = [
            ['缺失值填充数量', str(preprocessing_record.get('missing_count', 0))],
            ['异常值填充数量', str(preprocessing_record.get('outlier_count', 0))],
            ['对数转换', '是' if preprocessing_record.get('log_applied') else '否'],
            ['差分阶数', str(preprocessing_record.get('diff_order', 0))],
            ['最终记录数', str(preprocessing_record.get('final_length', '-'))]
        ]
        self._add_table(doc, ['处理步骤', '结果'], preprocess_rows, '表1-2 数据预处理情况')
        
        self._add_paragraph(doc, 
            f"数据预处理采用以下策略：对于缺失值，采用前值填充（Forward Fill）方法；"
            f"对于异常值，使用3σ法则检测后同样采用前值填充。"
            f"共处理{preprocessing_record.get('missing_count', 0)}个缺失值和"
            f"{preprocessing_record.get('outlier_count', 0)}个异常值。"
        )
        
        # 1.3 描述性统计
        self._add_heading(doc, '1.3 描述性统计', level=2)
        
        # ✅ 兼容单变量和多变量两种格式
        # 单变量: {'mean': 0.1, 'std': 0.2, ...}
        # 多变量: {'var1': {'mean': 0.1, ...}, 'var2': {'mean': 0.2, ...}}
        
        first_key = next(iter(descriptive_stats.keys()), None)
        is_multivariate = isinstance(descriptive_stats.get(first_key), dict)
        
        if is_multivariate:
            # 多变量：每个变量一个表格
            for var_name, var_stats in descriptive_stats.items():
                stats_rows = [
                    ['样本量', f"{var_stats.get('count', '-')}"],
                    ['均值', f"{var_stats.get('mean', 0):.4f}"],
                    ['标准差', f"{var_stats.get('std', 0):.4f}"],
                    ['最小值', f"{var_stats.get('min', 0):.4f}"],
                    ['最大值', f"{var_stats.get('max', 0):.4f}"],
                    ['偏度', f"{var_stats.get('skewness', 0):.4f}"],
                    ['峰度', f"{var_stats.get('kurtosis', 0):.4f}"]
                ]
                self._add_table(doc, ['统计量', '值'], stats_rows, f'表1-3 {var_name} 描述性统计')
        else:
            # 单变量：原有逻辑
            stats_rows = [
                ['样本量', f"{descriptive_stats.get('count', '-')}"],
                ['均值', f"{descriptive_stats.get('mean', 0):.4f}"],
                ['标准差', f"{descriptive_stats.get('std', 0):.4f}"],
                ['最小值', f"{descriptive_stats.get('min', 0):.4f}"],
                ['最大值', f"{descriptive_stats.get('max', 0):.4f}"],
                ['偏度', f"{descriptive_stats.get('skewness', 0):.4f}"],
                ['峰度', f"{descriptive_stats.get('kurtosis', 0):.4f}"]
            ]
            self._add_table(doc, ['统计量', '值'], stats_rows, '表1-3 描述性统计')
        
        # ========== 第2章：统计检验结果 ==========
        self._add_heading(doc, '2 统计检验结果', level=1)
        
        # 2.1-2.4 各项检验
        test_mapping = {
            'adf': ('2.1 单位根检验（ADF）', ['ADF统计量', 'p值', '结论']),
            'jb': ('2.2 正态性检验（JB）', ['JB统计量', 'p值', '结论']),
            'ljung_box': ('2.3 自相关检验（Ljung-Box Q）', ['Q统计量', 'p值', '滞后阶数', '结论']),
            'arch_lm': ('2.4 ARCH效应检验', ['LM统计量', 'p值', '滞后阶数', '结论'])
        }
        
        for test_key, (section_title, headers) in test_mapping.items():
            if test_key in test_results:
                self._add_heading(doc, section_title, level=2)
                result = test_results[test_key]
                
                if test_key == 'adf':
                    row = [
                        f"{result.get('test_statistic', '-'):.4f}",
                        f"{result.get('p_value', '-'):.4f}",
                        result.get('conclusion', '-')
                    ]
                elif test_key == 'jb':
                    row = [
                        f"{result.get('test_statistic', '-'):.4f}",
                        f"{result.get('p_value', '-'):.4f}",
                        result.get('conclusion', '-')
                    ]
                elif test_key in ['ljung_box', 'arch_lm']:
                    stat_key = 'test_statistic' if test_key == 'ljung_box' else 'lm_statistic'
                    p_key = 'p_value' if test_key == 'ljung_box' else 'lm_p_value'
                    row = [
                        f"{result.get(stat_key, '-'):.4f}",
                        f"{result.get(p_key, '-'):.4f}",
                        str(result.get('selected_lag', '-')),
                        result.get('conclusion', '-')
                    ]
                
                self._add_table(doc, headers, [row])
        
        # ========== 2.5-2.8 多变量检验（如果存在） ==========
        
        # 2.5 相关性矩阵
        if 'pearson_correlation' in test_results:
            self._add_heading(doc, '2.5 相关性分析', level=2)
            corr_result = test_results['pearson_correlation']
            corr_matrix = corr_result.get('correlation_matrix', {})
            
            if corr_matrix:
                # 构建相关系数矩阵表格
                variables = list(corr_matrix.keys())
                headers = ['变量'] + variables
                rows = []
                for var in variables:
                    row_data = [var]
                    for other_var in variables:
                        val = corr_matrix.get(var, {}).get(other_var, '-')
                        if isinstance(val, (int, float)):
                            row_data.append(f"{val:.4f}")
                        else:
                            row_data.append(str(val))
                    rows.append(row_data)
                
                self._add_table(doc, headers, rows, '表2-5 Pearson相关系数矩阵')
                
                # 高相关提示
                high_corr = corr_result.get('high_correlation_pairs', [])
                if high_corr:
                    pairs_text = "、".join([f"{p['var1']}-{p['var2']}({p['correlation']:.2f})" for p in high_corr[:5]])
                    self._add_paragraph(doc, f"注：相关系数绝对值大于0.7的变量对：{pairs_text}")
        
        # 2.6 VIF多重共线性检验
        if 'vif' in test_results:
            self._add_heading(doc, '2.6 多重共线性检验 (VIF)', level=2)
            vif_result = test_results['vif']
            vif_values = vif_result.get('vif_values', [])
            
            if vif_values:
                rows = []
                for item in vif_values:
                    var_name = item.get('variable', '-')
                    vif_val = item.get('vif', 0)
                    flag = '⚠ 共线' if item.get('has_collinearity') else '正常'
                    rows.append([var_name, f"{vif_val:.4f}", flag])
                
                self._add_table(doc, ['变量', 'VIF值', '状态'], rows, '表2-6 方差膨胀因子(VIF)检验结果')
                self._add_paragraph(doc, "注：VIF > 10 表示存在严重多重共线性。")
        
        # 2.7 Johansen协整检验
        if 'johansen' in test_results:
            self._add_heading(doc, '2.7 协整检验 (Johansen)', level=2)
            johansen_result = test_results['johansen']
            
            trace_test = johansen_result.get('trace_test', {})
            trace_stats = trace_test.get('statistics', [])
            trace_crit = trace_test.get('critical_values_95pct', [])
            
            if trace_stats and trace_crit:
                rows = []
                for i, (stat, crit) in enumerate(zip(trace_stats, trace_crit)):
                    conclusion = "拒绝H0" if stat > crit else "接受H0"
                    rows.append([f"r ≤ {i}", f"{stat:.4f}", f"{crit:.4f}", conclusion])
                
                self._add_table(doc, ['原假设', '迹统计量', '临界值(5%)', '结论'], rows, '表2-7 Johansen协整检验结果')
                
                rank = johansen_result.get('recommended_rank', 0)
                self._add_paragraph(doc, f"协整秩为 {rank}，{'表明变量间存在长期均衡关系' if rank > 0 else '未发现协整关系'}。")
        
        # 2.8 格兰杰因果检验
        if 'granger_causality' in test_results:
            self._add_heading(doc, '2.8 格兰杰因果检验', level=2)
            granger_result = test_results['granger_causality']
            causality_matrix = granger_result.get('causality_matrix', {})
            
            if causality_matrix:
                rows = []
                for pair_key, pair_info in causality_matrix.items():
                    if 'error' not in pair_info:
                        p_val = pair_info.get('p_value', 1)
                        lag = pair_info.get('optimal_lag', '-')
                        conclusion = "存在因果" if pair_info.get('is_granger_cause') else "不存在因果"
                        rows.append([pair_key, str(lag), f"{p_val:.4f}", conclusion])
                
                if rows:
                    self._add_table(doc, ['因果假设', '最优滞后', 'p值', '结论'], rows, '表2-8 格兰杰因果检验结果')
                    self._add_paragraph(doc, "注：格兰杰因果仅表示统计上的预测关系，不代表真实因果。")
        
        # ========== 第3章：模型构建与结果 ==========
        self._add_heading(doc, '3 模型构建与结果', level=1)
        
        # 3.1 模型选择依据
        self._add_heading(doc, '3.1 模型选择依据', level=2)
        self._add_paragraph(doc, 
            f"根据前述统计检验结果，本研究选择{model_type}模型进行时间序列建模分析。"
        )
        
        # 3.2 模型参数设定
        self._add_heading(doc, '3.2 模型参数设定', level=2)
        
        params = model_result.get('parameters', {})
        if model_type in ['ARIMA', 'ARMA']:
            order = model_result.get('selected_order', {})
            self._add_paragraph(doc, 
                f"模型阶数：p={order.get('p', '-')}, d={order.get('d', '-')}, q={order.get('q', '-')}"
            )
        elif model_type in ['GARCH', 'ARCH']:
            vol_eq = model_result.get('volatility_equation', {})
            mean_eq = model_result.get('mean_equation', {})
            dist = model_result.get('distribution', 'normal')
            self._add_paragraph(doc,
                f"波动率方程阶数：GARCH({vol_eq.get('p', '-')}, {vol_eq.get('q', '-')})\n"
                f"均值方程：AR({mean_eq.get('order', '-')})\n"
                f"分布假设：{dist}"
            )
        elif model_type in ['VAR', 'VECM']:
            self._add_paragraph(doc,
                f"滞后阶数：{model_result.get('selected_lag', model_result.get('lags', '-'))}"
            )
            if model_type == 'VECM':
                self._add_paragraph(doc,
                    f"协整秩：{model_result.get('cointegration_rank', '-')}"
                )
        
        # 3.3 参数估计结果
        self._add_heading(doc, '3.3 参数估计结果', level=2)
        
        all_params = params.get('all_params', params.get('coefficients', {}))
        p_values = params.get('p_values', {})
        
        param_rows = []
        for param_name, value in all_params.items():
            p_val = p_values.get(param_name, '-')
            if isinstance(p_val, float):
                p_val = f"{p_val:.4f}"
            param_rows.append([param_name, f"{value:.6f}" if isinstance(value, float) else str(value), str(p_val)])
        
        if param_rows:
            self._add_table(doc, ['参数', '估计值', 'p值'], param_rows, '表3-1 参数估计结果')
        
        # 3.4 拟合优度
        self._add_heading(doc, '3.4 拟合优度指标', level=2)
        metrics = model_result.get('metrics', {})
        metrics_rows = [
            ['R²', f"{metrics.get('r2', metrics.get('r2_average', '-')):.4f}" if metrics.get('r2') or metrics.get('r2_average') else '-'],
            ['AIC', f"{metrics.get('aic', '-'):.4f}" if metrics.get('aic') else '-'],
            ['BIC', f"{metrics.get('bic', '-'):.4f}" if metrics.get('bic') else '-'],
            ['Log-Likelihood', f"{metrics.get('log_likelihood', '-'):.4f}" if metrics.get('log_likelihood') else '-']
        ]
        self._add_table(doc, ['指标', '值'], metrics_rows, '表3-2 拟合优度指标')
        
        # 3.5 残差诊断
        self._add_heading(doc, '3.5 残差诊断', level=2)
        self._add_paragraph(doc, "模型残差的诊断检验结果详见图表部分的残差分析图。")
        
        # 3.6 经济学含义分析（AI生成）
        self._add_heading(doc, '3.6 经济学含义分析', level=2)
        
        # 调用Qwen生成分析
        economic_analysis = self.qwen_client.generate_economic_analysis(
            model_type=model_type,
            model_params=params,
            metrics=metrics,
            test_results=test_results
        )
        self._add_paragraph(doc, economic_analysis)
        
        # ========== 第4章：可视化图表 ==========
        self._add_heading(doc, '4 可视化图表', level=1)
        
        # 生成图表
        chart_paths = self.chart_generator.generate_all_charts(model_result, model_type)
        
        chart_titles = {
            'time_series': '图4-1 原始数据时间序列图',
            'acf': '图4-2 自相关函数图（ACF）',
            'pacf': '图4-3 偏自相关函数图（PACF）',
            'residuals_series': '图4-4 残差时间序列图',
            'residuals_hist': '图4-5 残差直方图',
            'qq_plot': '图4-6 残差Q-Q图',
            'fitted_vs_actual': '图4-7 拟合值与实际值对比图',
            'conditional_volatility': '图4-8 条件波动率图',
            'impulse_response': '图4-9 脉冲响应函数图',
            'variance_decomposition': '图4-10 方差分解图'
        }
        
        for chart_key, chart_path in chart_paths.items():
            caption = chart_titles.get(chart_key, chart_key)
            self._add_image(doc, chart_path, caption)
        
        # 保存报告到report.docx（固定文件名）
        report_path = os.path.join(self.report_dir, "report.docx")
        doc.save(report_path)
        logger.info(f"报告已生成: {report_path}")

        # 保存报告元数据到meta.json
        import json
        import time
        meta_path = os.path.join(self.report_dir, "meta.json")
        meta_data = {
            "model_type": model_type,
            "data_source": data_source_info.get('name', '-'),
            "created_at": int(time.time()),
            "report_name": f"{data_source_info.get('name', '数据')}_{model_type}模型",
            "status": "active",
            "metrics": {
                "aic": model_result.get('metrics', {}).get('aic'),
                "bic": model_result.get('metrics', {}).get('bic'),
                "r2": model_result.get('metrics', {}).get('r2') or model_result.get('metrics', {}).get('r2_average')
            }
        }
        with open(meta_path, 'w', encoding='utf-8') as f:
            json.dump(meta_data, f, ensure_ascii=False, indent=2)

        logger.info(f"报告元数据已保存: {meta_path}")

        return report_path
